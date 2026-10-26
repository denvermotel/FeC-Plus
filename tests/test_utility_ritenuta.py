#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# FeC-Plus - v0.04 dev
"""Elenco Excel di Utility con le colonne della ritenuta d'acconto."""

import os
import shutil
import sys
import tempfile
import unittest
from unittest.mock import patch

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import fec_utility
from _finti import Risposta

try:
    import openpyxl
except ImportError:
    openpyxl = None

CON = (b'<?xml version="1.0"?><FatturaElettronica><FatturaElettronicaBody><DatiGenerali>'
       b'<DatiGeneraliDocumento><DatiRitenuta><TipoRitenuta>RT01</TipoRitenuta>'
       b'<ImportoRitenuta>200.00</ImportoRitenuta><AliquotaRitenuta>20.00</AliquotaRitenuta>'
       b'</DatiRitenuta></DatiGeneraliDocumento></DatiGenerali></FatturaElettronicaBody>'
       b'</FatturaElettronica>')
SENZA = b'<?xml version="1.0"?><FatturaElettronica><Numero>1</Numero></FatturaElettronica>'

VOCI = [
    {"tipoInvio": "FPR", "idFattura": "1", "numeroFattura": "1", "dataFattura": "10/01/2026",
     "tipoDocumento": "TD06", "denominazioneCliente": "Rossi Mario"},
    {"tipoInvio": "FPR", "idFattura": "2", "numeroFattura": "2", "dataFattura": "11/01/2026",
     "tipoDocumento": "TD01", "denominazioneCliente": "Verdi Luigi"},
    {"tipoInvio": "FPR", "idFattura": "3", "numeroFattura": "3", "dataFattura": "12/01/2026",
     "tipoDocumento": "TD04", "denominazioneCliente": "Rossi Mario"},
]
FILE = {"FPR1": CON, "FPR2": SENZA, "FPR3": CON}


class Sessione:
    def __init__(self):
        self.url = []

    def get(self, url, **kwargs):
        self.url.append(url)
        chiave = next(k for k in FILE if f"/{k}?" in url)
        if FILE[chiave] is None:
            return Risposta(status=500)
        return Risposta(contenuto=FILE[chiave], filename=f"{chiave}.xml")


class Auth:
    def __init__(self):
        self.session = Sessione()
        self.headers = {}


@unittest.skipIf(openpyxl is None, "richiede openpyxl")
class TestExcelRitenuta(unittest.TestCase):

    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.tmp, True)
        patch.object(fec_utility, "_lista_fatture", return_value=[dict(v) for v in VOCI]).start()
        patch.object(fec_utility, "_dettaglio_fattura", return_value=None).start()
        self.addCleanup(patch.stopall)

    def _excel(self, **extra):
        self.auth = Auth()
        percorso = fec_utility.elenco_fatture_excel(
            self.auth, "RSSMRA80A01H501U", "01234567890", "emesse", "01012026", "31012026",
            dest_dir=self.tmp, sottocartella=False, log=lambda _m: None, **extra)
        foglio = openpyxl.load_workbook(percorso).active
        righe = list(foglio.iter_rows(values_only=True))
        return foglio, righe[0], righe[1:]

    def test_colonne_e_valori(self):
        foglio, testa, righe = self._excel(ritenuta=True)
        imp, tipo, perc = (testa.index("Ritenuta importo"), testa.index("Ritenuta tipo"),
                           testa.index("Ritenuta %"))
        self.assertEqual([(r[tipo], r[perc], r[imp]) for r in righe[:3]],
                         [("RT01", "20", 200.0), (None, None, None), ("RT01", "20", -200.0)])
        # riga TOTALI: le note di credito si sottraggono
        self.assertEqual(righe[3][0], "TOTALI")
        self.assertEqual(righe[3][imp], 0.0)

    def test_righe_con_ritenuta_evidenziate(self):
        foglio, testa, righe = self._excel(ritenuta=True)
        colori = [foglio.cell(row=i, column=1).fill.start_color.rgb for i in (2, 3, 4)]
        self.assertEqual(colori[0], colori[2])
        self.assertNotEqual(colori[0], colori[1])

    def test_senza_opzione_nessuna_colonna_e_nessun_download(self):
        foglio, testa, righe = self._excel()
        self.assertNotIn("Ritenuta importo", testa)
        self.assertEqual(self.auth.session.url, [])

    def test_file_non_scaricato_distinto_da_non_leggibile(self):
        FILE["FPR2"] = None
        self.addCleanup(FILE.__setitem__, "FPR2", SENZA)
        foglio, testa, righe = self._excel(ritenuta=True)
        self.assertEqual(righe[1][testa.index("Ritenuta tipo")], "file non scaricato")

    def test_in_predisposizione_nessun_download(self):
        voci = [dict(v) for v in VOCI]
        voci[1]["fileDownload"] = {"fileDownload": 2}
        with patch.object(fec_utility, "_lista_fatture", return_value=voci):
            foglio, testa, righe = self._excel(ritenuta=True)
        self.assertFalse(any("/FPR2?" in u for u in self.auth.session.url))
        self.assertEqual(righe[1][testa.index("Ritenuta tipo")], "file non ancora disponibile")

    def test_righe_non_leggibili_con_un_altro_colore(self):
        FILE["FPR2"] = b"\x30\x82\x00\x10non una busta"
        self.addCleanup(FILE.__setitem__, "FPR2", SENZA)
        foglio, testa, righe = self._excel(ritenuta=True)
        con, non_leggibile = (foglio.cell(row=i, column=1).fill.start_color.rgb
                              for i in (2, 3))
        self.assertNotEqual(con, non_leggibile)
        self.assertNotIn(non_leggibile, (None, "00000000"))

    def test_file_illeggibile_segnalato_in_tabella(self):
        FILE["FPR2"] = b"\x30\x82\x00\x10non una busta"
        self.addCleanup(FILE.__setitem__, "FPR2", SENZA)
        foglio, testa, righe = self._excel(ritenuta=True)
        self.assertEqual(righe[1][testa.index("Ritenuta tipo")], "non leggibile")


if __name__ == "__main__":
    unittest.main()
