#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# FeC-Plus - v0.04 dev
"""Export Excel dei corrispettivi su sessione finta (nessuna rete)."""

import os
import shutil
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import fec_corrispettivi as fc
import fec_utility
from _finti import Auth, SessioneRotte

try:
    import openpyxl
except ImportError:              # dipendenza opzionale
    openpyxl = None

_SINTESI = {"registratoriInvii": "1", "multicassaInvii": "0", "distributoriInvii": "2",
            "carburantiInvii": "0", "documentiCommerciali": "1", "registratoriCassa": "0",
            "torretteEnergia": "0"}


def _rotte():
    return [
        ("/sintesi/dal/", _SINTESI),
        ("tipoCorrispettivo/RT", {"corrispettivi": [{
            "idInvio": "1", "timeRilevazione": "10/07/2026 20:00",
            "matricolaDispositivo": "M1", "dettaglio": True}], "totaleCorrispettivi": 1}),
        ("/dettaglio/RT1", {"datiContabiliRT_MC": [
            {"aliquota": "22.00", "imponibile": "+000000000100,00",
             "imposta": "+000000000022,00", "natura": "", "ventilazione": ""}]}),
        # letture precedenti dei distributori: finestra dei 3 mesi prima del 01/07
        ("/dal/01042026/al/30062026/piva/01234567890/tipoCorrispettivo/DA", {
            "corrispettivi": [{"idInvio": "5", "timeRilevazione": "20/06/2026 10:00",
                               "matricolaDispositivo": "D1", "modalitaRilevazione": "CUMULATO",
                               "importo": "1000,00"}], "totaleCorrispettivi": 1}),
        ("tipoCorrispettivo/DA", {"corrispettivi": [
            {"idInvio": "6", "timeRilevazione": "05/07/2026 10:00",
             "matricolaDispositivo": "D1", "modalitaRilevazione": "CUMULATO",
             "importo": "1100,00", "dettaglio": False},
            {"idInvio": "7", "timeRilevazione": "25/07/2026 10:00",
             "matricolaDispositivo": "D1", "modalitaRilevazione": "CUMULATO",
             "importo": "1250,00", "dettaglio": False}], "totaleCorrispettivi": 2}),
        ("/dc/ricerca/", {"elenco": [{"id": "77", "dataEmissione": "2026-07-03"}],
                          "totaleElementi": 1}),
        ("/dc/dettaglio/77", {"listaAliquota": [
            {"aliquotaIva": 10, "natura": "", "imponibile": 50, "imposta": 5}]}),
    ]


@unittest.skipIf(openpyxl is None, "richiede openpyxl")
class TestExcelCorrispettivi(unittest.TestCase):

    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.tmp, True)

    def _esegui(self, **kw):
        return fc.elenco_corrispettivi_excel(
            Auth(SessioneRotte(_rotte())), "CFCLIENTE", "01234567890",
            "01072026", "31072026", dest_dir=self.tmp, sottocartella=False,
            log=lambda *_: None, **kw)

    def test_un_file_per_canale_e_dispositivo(self):
        esito = self._esegui()
        nomi = sorted(os.path.basename(p) for p in esito.percorsi)
        self.assertEqual(nomi, [
            "01234567890_01072026-31072026_corrispettivi_M1.xlsx",
            "01234567890_01072026-31072026_distributori_D1.xlsx",
            "01234567890_01072026-31072026_documenti_commerciali.xlsx"])
        self.assertEqual(esito.avvisi, [])

    def test_distributori_per_differenza_e_totale_del_solo_venduto(self):
        esito = self._esegui()
        percorso = next(p for p in esito.percorsi if "distributori" in p)
        ws = openpyxl.load_workbook(percorso).active
        intest = [c.value for c in ws[1]]
        col = {nome: i for i, nome in enumerate(intest)}
        righe = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
        self.assertEqual(righe[0][col["Lettura precedente"]], 1000.0)
        self.assertEqual(righe[0][col["Venduto dell'intervallo"]], 100.0)
        self.assertEqual(righe[1][col["Venduto dell'intervallo"]], 150.0)
        totali = righe[-1]
        self.assertEqual(totali[0], "TOTALI")
        self.assertEqual(totali[col["Venduto dell'intervallo"]], 250.0)
        self.assertIn(totali[col["Totale venduto (contatore)"]], ("", None))
        self.assertIn(totali[col["Lettura precedente"]], ("", None))

    def test_distributori_dati_grezzi(self):
        esito = self._esegui(modalita_da="grezzi")
        percorso = next(p for p in esito.percorsi if "distributori" in p)
        intest = [c.value for c in openpyxl.load_workbook(percorso).active[1]]
        self.assertIn("Venduto", intest)
        self.assertNotIn("Venduto dell'intervallo", intest)

    def test_modalita_sconosciuta(self):
        with self.assertRaises(fc.DownloadError):
            self._esegui(modalita_da="boh")

    def test_piva_obbligatoria(self):
        with self.assertRaises(fc.DownloadError):
            fc.elenco_corrispettivi_excel(Auth(SessioneRotte(_rotte())), "CF", "",
                                          "01072026", "31072026", dest_dir=self.tmp,
                                          log=lambda *_: None)

    def test_nessun_invio(self):
        vuota = [("/sintesi/dal/", {k: "0" for k in _SINTESI})]
        with self.assertRaises(fec_utility.NessunDato):
            fc.elenco_corrispettivi_excel(Auth(SessioneRotte(vuota)), "CF", "P",
                                          "01072026", "31072026", dest_dir=self.tmp,
                                          sottocartella=False, log=lambda *_: None)
        self.assertEqual(os.listdir(self.tmp), [])


class TestDump(unittest.TestCase):

    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.tmp, True)

    def test_sintesi_elenco_e_dettaglio_dei_tipi_presenti(self):
        percorsi = fc.dump_json(Auth(SessioneRotte(_rotte())), "01234567890",
                                "01072026", "31072026", self.tmp, log=lambda *_: None)
        nomi = sorted(os.path.basename(p) for p in percorsi)
        self.assertIn("dump_corr_sintesi.json", nomi)
        self.assertIn("dump_corr_elenco_RT.json", nomi)
        self.assertIn("dump_corr_dettaglio_RT.json", nomi)
        self.assertIn("dump_corr_elenco_DC.json", nomi)
        self.assertNotIn("dump_corr_elenco_MC.json", nomi)


if __name__ == "__main__":
    unittest.main()
